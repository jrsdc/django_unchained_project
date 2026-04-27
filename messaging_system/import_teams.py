import os
import re
from pathlib import Path
from collections import defaultdict

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "messaging_system.settings")

import django
django.setup()

from openpyxl import load_workbook
from teams.models import Team


BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "team_registry.xlsx"


def normalise(text):
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


def clean(value):
    if value is None:
        return ""
    value = str(value).strip()
    if value in ["#REF!", "#NAME?", "nan", "None"]:
        return ""
    return value


def make_email(team_name):
    slug = re.sub(r"[^a-z0-9]+", "", team_name.lower())
    return f"{slug}@sky.com"


def fix_url(value):
    value = clean(value)
    if not value:
        return ""
    if not value.startswith(("http://", "https://")):
        value = "https://" + value
    return value


def split_dependencies(value):
    value = clean(value)
    if not value or value.lower() == "none":
        return []
    return [
        item.strip()
        for item in re.split(r"[,;\n]+", value)
        if item.strip()
    ]


if not EXCEL_FILE.exists():
    print("ERROR: team_registry.xlsx not found.")
    print("Put the Excel file in the same folder as manage.py and rename it to team_registry.xlsx")
    raise SystemExit


workbook = load_workbook(EXCEL_FILE, data_only=True)
sheet = workbook.active

headers = [clean(cell.value) for cell in sheet[1]]
header_map = {normalise(header): index for index, header in enumerate(headers)}


def get(row, *possible_headers):
    for header in possible_headers:
        index = header_map.get(normalise(header))
        if index is not None and index < len(row):
            return clean(row[index])
    return ""


rows = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    team_name = get(row, "Team Name")

    if not team_name:
        continue

    department = get(row, "Department") or "Unknown Department"
    manager = get(row, "Team Leader") or "Unknown Manager"
    team_type = get(row, "Jira Project Name") or department
    description = get(row, "Development Focus Areas") or "No description provided."

    skills = (
        get(row, "Key Skills & Technologies")
        or get(row, "Development Focus Areas")
        or "General engineering"
    )

    repository = fix_url(
        get(
            row,
            "Project (codebase) (Github Repo Link)",
            "Project codebase Github Repo",
            "Project codebase",
            "Github Repo",
            "GitHub Repo",
        )
    )

    downstream = get(row, "Downstream Dependencies") or "None"

    rows.append({
        "name": team_name,
        "team_type": team_type,
        "department": department,
        "manager": manager,
        "description": description,
        "skills": skills,
        "downstream_dependencies": downstream,
        "email": make_email(team_name),
        "repository": repository,
    })


# Build upstream dependencies automatically by reversing downstream links
upstream_map = defaultdict(set)

for item in rows:
    source_team = item["name"]
    for downstream_team in split_dependencies(item["downstream_dependencies"]):
        upstream_map[downstream_team].add(source_team)


created_count = 0
updated_count = 0

for item in rows:
    upstream = ", ".join(sorted(upstream_map.get(item["name"], []))) or "None"

    team, created = Team.objects.update_or_create(
        name=item["name"],
        defaults={
            "team_type": item["team_type"],
            "department": item["department"],
            "manager": item["manager"],
            "description": item["description"],
            "skills": item["skills"],
            "upstream_dependencies": upstream,
            "downstream_dependencies": item["downstream_dependencies"] or "None",
            "email": item["email"],
            "repository": item["repository"],
        }
    )

    if created:
        created_count += 1
    else:
        updated_count += 1


print(f"Import complete.")
print(f"Created: {created_count}")
print(f"Updated: {updated_count}")
print(f"Total rows imported: {len(rows)}")